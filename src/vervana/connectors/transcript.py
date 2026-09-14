"""Transcript / video-quote ingest (M6, Part 4.2) — MANUAL-FILE path only.

Gate 1 forbids code that *fetches* from YouTube or a parser that consumes *raw* YouTube
transcripts. This connector does neither: it ingests a HUMAN-CURATED spreadsheet export
(a person watched the videos and transcribed the price statements), exactly like the
quick-commerce manual panel. There is no network fetch and no raw-transcript parsing here.
The connector is pluggable and can be permanently dropped without touching anything else.

Video quotes are `source_class = quote_indicative` and are NEVER blended with executed
prices (the repository guard enforces this). The unit is usually unstated, so the
canonical ₹/kg stays NULL rather than being guessed (Part 3.3).
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from vervana.connectors.base import Connector, IngestResult
from vervana.db.base import CanonicalType, SourceClass, TimeBasis
from vervana.models.entities import Market
from vervana.money import rupees_to_paise
from vervana.repository.prices import insert_observation
from vervana.repository.registry import resolve_by_name
from vervana.time import IST

# Column headers in the curated corpus.
C_DATE, C_DAY, C_CHANNEL, C_COMMODITY, C_VARIETY = (
    "Date",
    "Day",
    "Channel / Mandi",
    "Commodity",
    "Variety / Origin (as quoted)",
)
C_LOW, C_HIGH, C_UNIT, C_ARRIVALS, C_CARRYOVER = (
    "Price Low",
    "Price High",
    "Unit",
    "Arrivals (trucks)",
    "Carryover",
)
C_COMMENTARY, C_QUOTE, C_VIDEO, C_FLAG = (
    "Market Commentary",
    "Raw Quote (transcript)",
    "Source Video",
    "Flag",
)

_KG_UNITS = {"kg", "kilo", "kilogram", "रु/किलो", "rs/kg"}

# --------------------------------------------------------------------------------------
# Ground-proof commodity labels (curated-workbook "Commodity" column).
#
# The vloggers often quote a combined label for one price range ("Tomato + Peas
# 1000-1200"). When EVERY part of the label names a specific registry commodity we
# record the same quoted range against each component, flagged `combined_quote:<label>`
# so it is never mistaken for a per-commodity-only quote (Part 4.2: flag, never guess).
# Keys are normalize()d forms (lowercase, punctuation stripped).
# --------------------------------------------------------------------------------------
_COMBINED_LABEL_COMPONENTS: dict[str, tuple[str, ...]] = {
    "tomato peas": ("Tomato", "Green Peas"),          # matar
    "chilli capsicum": ("Green Chilli", "Capsicum"),  # hari mirch + shimla mirch
    "apple capsicum": ("Apple", "Capsicum"),
    "mosambi guava": ("Mosambi", "Guava"),
}

# Labels that can never resolve to one commodity. They stay rejected, each with an
# explicit reason so coverage reporting shows WHY the rows are out (Part 3.2: a tie or
# an unknown is unresolved, never a guess).
_GENERIC_BUCKETS = {
    "fruits all",
    "vegetables all",
    "fruits all vegetables all",
    "vegetables all fruits all",
    "other",
}
# Precious-metals commentary in the same vlogs - not produce, out of platform scope.
_NON_PRODUCE_LABELS = {"gold silver"}


def _combined_components(label: str) -> tuple[str, ...] | None:
    """Return the specific commodities a combined label covers, else None."""
    from vervana.matching.normalize import normalize

    return _COMBINED_LABEL_COMPONENTS.get(normalize(label))


def _unresolved_reason(label: str) -> str:
    """Classify an unresolvable label for honest coverage/reject reporting."""
    from vervana.matching.normalize import normalize

    n = normalize(label)
    if n in _NON_PRODUCE_LABELS:
        return "non_produce_label"
    if n in _GENERIC_BUCKETS:
        return "generic_bucket_label"
    if "fruits all" in n or "vegetables all" in n:
        return "combined_with_generic_bucket"
    return "unresolved_commodity"


class TranscriptConnector(Connector):
    name = "transcript"

    def fetch_raw(self, **params) -> list[dict]:
        # RISK[R8-YOUTUBE-TOS]: This corpus is YouTube-derived. It cannot be both a legal
        # risk and a data-room asset (R8), and Gate 1's written opinion on commercial use
        # of extracted price statements has not been confirmed. There is deliberately NO
        # fetch path — data enters only via a human-curated file — and the raw corpus is
        # kept out of git (data/raw is git-ignored) so it cannot silently become a
        # "shipped asset". If the legal opinion comes back negative, drop this connector.
        # Evidence: docs/RISK_REGISTER.md#r8-youtube-tos; Gate 1
        # Verdict: PENDING — legal opinion not yet confirmed; treated as manual import only.
        raise NotImplementedError(
            "transcript ingest has no fetch path (Gate 1: no YouTube fetching). "
            "Use import_csv with a human-curated corpus export."
        )

    def import_csv(self, session, path: str | Path) -> IngestResult:
        with Path(path).open(encoding="utf-8-sig") as fh:
            return self.ingest(session, list(csv.DictReader(fh)), mode="csv")

    @staticmethod
    def records_from_upload(filename: str, content: bytes) -> list[dict]:
        """Parse the repeatable ground-proof CSV/XLSX format in memory."""
        suffix = Path(filename or "").suffix.lower()
        if suffix == ".csv":
            return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = book.active.iter_rows(values_only=True)
            try:
                headers = [str(v).strip() if v is not None else "" for v in next(rows)]
            except StopIteration:
                return []
            return [dict(zip(headers, row, strict=False)) for row in rows if any(v is not None for v in row)]
        raise ValueError("unsupported file type: send .xlsx, .xlsm, or .csv")

    def ingest(self, session, records: list[dict], *, mode: str = "csv") -> IngestResult:
        result = IngestResult(rows_in=len(records))
        for rec in records:
            market_id = self._resolve_market(session, rec.get(C_CHANNEL, ""))
            if market_id is None:
                result.reject(f"unresolved_market: {rec.get(C_CHANNEL)}", rec)
                continue
            label = (rec.get(C_COMMODITY) or "").strip()
            components = _combined_components(label)
            if components is not None:
                component_ids = []
                for cname in components:
                    cid = resolve_by_name(
                        session, name=cname, canonical_type=CanonicalType.commodity
                    )
                    if cid is None:
                        component_ids = None
                        break
                    component_ids.append(cid)
                if component_ids is None:
                    result.reject(f"combined_component_unresolved: {label}", rec)
                    continue
                for cid in component_ids:
                    self._ingest_one(
                        session,
                        rec,
                        market_id=market_id,
                        commodity_id=cid,
                        extra_flags=[f"combined_quote:{label}"],
                        result=result,
                    )
                continue

            commodity_id = resolve_by_name(
                session,
                name=label,
                canonical_type=CanonicalType.commodity,
            )
            if commodity_id is None:
                result.reject(f"{_unresolved_reason(label)}: {label}", rec)
                continue
            self._ingest_one(
                session,
                rec,
                market_id=market_id,
                commodity_id=commodity_id,
                extra_flags=[],
                result=result,
            )
        return result

    def _ingest_one(
        self,
        session,
        rec: dict,
        *,
        market_id: int,
        commodity_id: int,
        extra_flags: list[str],
        result: IngestResult,
    ) -> None:
        """Ingest one row against one resolved commodity.

        Combined labels ("Tomato + Peas") call this once per component, so one
        workbook row can produce one observation per specific commodity, each
        carrying the `combined_quote` flag. Per-commodity dedupe is unchanged.
        """

        low = self._price(rec.get(C_LOW))
        high = self._price(rec.get(C_HIGH))
        if low is None and high is None:
            result.reject("no_price_extracted", rec)
            return

        flags = self._flags(rec, low, high) + list(extra_flags)
        # RISK[R3-VIDEO-PROVENANCE]: These are YouTube-quoted ranges treated as prices.
        # UNVERIFIED against executed trades — the ground-truth study (M5) compares them
        # to trader invoices. Until then every row here is a SENTIMENT/quote signal, not
        # an executed price; source_class=quote_indicative keeps them un-blendable with
        # executed prices, and the unit is usually unstated so canonical ₹/kg stays NULL.
        # Evidence: docs/RISK_REGISTER.md#r3-video-provenance; assumptions A3
        # Verdict: PENDING
        lo = low if low is not None else high
        hi = high if high is not None else low
        if lo > hi:
            result.reject("range_disorder", rec)
            return

        unit_raw = (rec.get(C_UNIT) or "unstated").strip() or "unstated"
        canonical = None
        if unit_raw.lower() in _KG_UNITS:
            canonical = round((lo + hi) / 2)

        source_url = (rec.get(C_VIDEO) or "").strip()
        raw_quote = (rec.get(C_QUOTE) or "").strip()
        if not source_url:
            result.reject("missing_source_url", rec)
            return
        if not raw_quote:
            result.reject("missing_raw_quote", rec)
            return
        evidence_json = json.dumps(
            {
                "source_type": "youtube_ground_proof",
                "quote": raw_quote,
                "channel_mandi": rec.get(C_CHANNEL),
                "variety": rec.get(C_VARIETY),
                "commentary": rec.get(C_COMMENTARY),
                "arrivals": rec.get(C_ARRIVALS),
                "carryover": rec.get(C_CARRYOVER),
                "source_flag": rec.get(C_FLAG),
                "detected_flags": flags,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        observed_at = self._date(rec.get(C_DATE))
        from sqlalchemy import select
        from vervana.models.observations import PriceObservation
        duplicate = session.scalar(
            select(PriceObservation.id).where(
                PriceObservation.commodity_id == commodity_id,
                PriceObservation.market_id == market_id,
                PriceObservation.source_class == SourceClass.quote_indicative,
                PriceObservation.observed_at == observed_at,
                PriceObservation.source_url == source_url,
                PriceObservation.raw_quote == evidence_json,
            ).limit(1)
        )
        if duplicate is not None:
            result.reject("duplicate_already_ingested", rec)
            return

        insert_observation(
            session,
            commodity_id=commodity_id,
            market_id=market_id,
            source_class=SourceClass.quote_indicative,
            price_low_paise=lo,
            price_high_paise=hi,
            price_point_paise=lo if low is not None and high is None else None,
            unit_raw=unit_raw,
            canonical_price_paise_per_kg=canonical,
            unit_conversion_confidence=1.0 if canonical is not None else None,
            source_url=source_url,
            raw_quote=evidence_json,
            observed_at=observed_at,
            time_basis=TimeBasis.single_daily_quote,
        )
        result.accept()

    # --- helpers ---------------------------------------------------------------
    @staticmethod
    def _resolve_market(session, channel: str):
        text = (channel or "").lower()
        from sqlalchemy import select

        for m in session.scalars(select(Market)):
            if m.canonical_name.lower() in text:
                return m.id
        return None

    @staticmethod
    def _price(value) -> int | None:
        if value in (None, "", "None"):
            return None
        try:
            # Numbers exports numeric cells; CSV re-read gives strings like "15.0".
            return rupees_to_paise(str(int(float(value))))
        except Exception:
            return None

    @staticmethod
    def _flags(rec: dict, low: int | None, high: int | None) -> list[str]:
        flags: list[str] = []
        src = (rec.get(C_FLAG) or "").lower()
        if "merged digits" in src:
            flags.append("source_flag:possible_merged_digits")
        # Our own digit-merge detection: a 4-digit single price that splits into a plausible
        # tight range, e.g. 4546 -> 45–46. FLAG ONLY, never silently correct (Part 4.2).
        for label, val in (("low", rec.get(C_LOW)), ("high", rec.get(C_HIGH))):
            s = re.sub(r"\.0$", "", str(val or "").strip())
            if re.fullmatch(r"\d{4}", s):
                a, b = int(s[:2]), int(s[2:])
                if 0 < b - a <= 3 and 5 <= a <= 99:
                    flags.append(f"digit_merge_suspected:{label} {s}->{a}-{b}")
        # Range-sanity flags (paise). An unstated-unit fresh-produce quote above ~₹5000 or a
        # range spanning >10x is almost certainly a transcript artifact or a unit mix-up.
        for label, v in (("low", low), ("high", high)):
            if v is not None and v > 500_000:  # > ₹5000
                flags.append(f"price_implausibly_high:{label} ₹{v // 100}")
        if low and high and high > low * 10:
            flags.append(f"range_implausible_ratio:{high // 100}/{low // 100}")
        return flags

    @staticmethod
    def _date(value):
        from datetime import datetime

        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                d = datetime.strptime((value or "").strip(), fmt)
                return datetime(d.year, d.month, d.day, tzinfo=IST)
            except ValueError:
                continue
        from vervana.time import now_utc

        return now_utc()
